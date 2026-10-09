# Builds the Meta Commerce Manager product feed (CSV, TSV, XLSX) for KIKFIA.
# Usage: python3 build_feed.py <commit-sha-hosting-the-images>
# Prices are left blank on purpose: KIKFIA must enter its own prices before uploading.
import csv, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

SHA = sys.argv[1]
IMG = 'https://raw.githubusercontent.com/stkvideos3-dot/KIKFIA/%s/meta-catalog/images/%%s.jpg' % SHA
LINK = 'https://kikfia.com/'
TAIL = (' Factory-built and customized to your site: choose the layout, size, finishes and colors. '
        'Price shown is a starting price; final price depends on options, quantity and delivery location. '
        'Contact KIKFIA for a free consultation and a written quote.')

P = [
 ('KF-ADU-STUDIO', 'KIKFIA Backyard Studio ADU', 'ADU Houses',
  'A modern backyard studio with full-height glass doors, an insulated shell and a roof ready for solar panels. Ideal as a home office, creative studio, gym or guest room.',
  'adu-backyard-studio', ['interior-bedroom', 'interior-kitchenette']),
 ('KF-ADU-VERANDA', 'KIKFIA ADU House with Covered Veranda', 'ADU Houses',
  'A complete second home with kitchen, bathroom and bedrooms plus a covered front veranda. Built for family members, caregivers or long-term rental income.',
  'adu-veranda-house', ['interior-kitchenette', 'interior-bedroom']),
 ('KF-EXP-20', 'KIKFIA Expandable House 20 ft', 'Expandable Houses',
  'Ships as a compact unit and unfolds on site into a finished home of about 370 sq ft. Studio to 2-bedroom layouts with kitchen, bathroom and pre-installed wiring and plumbing.',
  'expandable-20ft', ['expandable-resort-row', 'interior-bedroom', 'factory-finishing']),
 ('KF-EXP-40', 'KIKFIA Expandable House 40 ft', 'Expandable Houses',
  'Unfolds on site into a finished home of about 745 sq ft with up to 3 bedrooms, a full kitchen and bathroom. Glass-front, veranda and cladding options available.',
  'expandable-40ft', ['expandable-veranda-row', 'expandable-village', 'interior-kitchenette']),
 ('KF-TINY-CAPSULE', 'KIKFIA Capsule Tiny House', 'Tiny Houses',
  'A rounded, fully insulated capsule cabin with large windows, built for year-round comfort. Popular for glamping sites, vacation rentals and backyard retreats.',
  'capsule-tiny-house', ['capsule-resort-aerial']),
 ('KF-TINY-CONTAINER', 'KIKFIA Container Tiny House', 'Tiny Houses',
  'A container-style tiny house with sliding glass doors, a compact kitchen and a full bathroom. Easy to move and ideal for first homes and short-term rentals.',
  'container-tiny-house', ['container-kitchen', 'interior-bedroom']),
 ('KF-MOD-HOME', 'KIKFIA Two-Story Modular Home', 'Modular Homes',
  'A spacious family home built from factory-made modules, with balconies, large windows and your choice of facade finishes.',
  'modular-two-story-home', ['site-units-stacked']),
 ('KF-MOD-VILLA', 'KIKFIA Modern Modular Villa', 'Modular Homes',
  'A single-story modular villa with floor-to-ceiling glass, a wide deck and open-plan living. Ideal for luxury homes and vacation properties.',
  'modular-glass-villa', ['interior-kitchenette']),
 ('KF-GUEST-ROOF', 'KIKFIA Guest House with Rooftop Terrace', 'Guest Houses',
  'A private guest suite with bedroom, bathroom and kitchenette, a wraparound deck and an optional rooftop terrace with stairs.',
  'guest-house-rooftop', ['guest-house-front']),
 ('KF-OFFICE', 'KIKFIA Modular Office Building', 'Portable Offices',
  'Ready-to-use offices from a single unit to two-story buildings, with glass entries, lighting and power. Ideal for job sites, sales offices and growing teams.',
  'modular-office-building', ['office-single-story', 'office-interior']),
 ('KF-RESORT', 'KIKFIA Resort Villas and Hotel Units', 'Resorts and Hospitality',
  'Resort villas, hotel units and glamping cabins that can open in phases and grow with demand. Decks, verandas and themed finishes available.',
  'resort-villas', ['capsule-resort-aerial', 'expandable-resort-row']),
 ('KF-COMMERCIAL', 'KIKFIA Commercial Container Units', 'Commercial Buildings',
  'Stackable container units for cafés, shops, showrooms and hotel rooms, delivered with your colors, glazing and branding.',
  'commercial-container-units', ['cafe-pavilion']),
 ('KF-SITE-UNITS', 'KIKFIA Storage and Site Units', 'Storage and Industrial',
  'Secure, weather-resistant units for storage, workshops and site facilities. Join side by side or stack up to three levels.',
  'site-units-stacked', ['storage-unit']),
 ('KF-CLINIC-SCHOOL', 'KIKFIA Modular Clinic and Classroom Buildings', 'Schools and Clinics',
  'Modular buildings for classrooms, training centers, clinics and staff housing, delivered on tight schedules.',
  'clinic-school-building', ['training-room']),
]

COLS = ['id', 'title', 'description', 'availability', 'condition', 'price', 'link', 'image_link',
        'additional_image_link', 'brand', 'product_type', 'custom_label_0']
rows = []
for pid, title, ptype, desc, main, extra in P:
    rows.append({
        'id': pid, 'title': title, 'description': desc + TAIL,
        'availability': 'available for order', 'condition': 'new', 'price': '',
        'link': LINK, 'image_link': IMG % main,
        'additional_image_link': ','.join(IMG % e for e in extra),
        'brand': 'KIKFIA', 'product_type': 'Portable Houses > ' + ptype,
        'custom_label_0': 'Core product' if ptype in ('ADU Houses', 'Expandable Houses', 'Tiny Houses', 'Modular Homes') else 'More solutions',
    })

for name, delim in (('kikfia-meta-catalog.csv', ','), ('kikfia-meta-catalog.tsv', '\t')):
    with open(name, 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, fieldnames=COLS, delimiter=delim)
        w.writeheader(); w.writerows(rows)

wb = Workbook(); ws = wb.active; ws.title = 'Products'
ws.append(COLS)
for r in rows: ws.append([r[c] for c in COLS])
head = Font(bold=True, color='FFFFFF'); fill = PatternFill('solid', fgColor='15302A')
for c in ws[1]: c.font = head; c.fill = fill
yellow = PatternFill('solid', fgColor='FFF2A8')
for row in ws.iter_rows(min_row=2, min_col=6, max_col=6):
    for c in row: c.fill = yellow
widths = {'A': 18, 'B': 44, 'C': 70, 'D': 20, 'E': 10, 'F': 16, 'G': 22, 'H': 60, 'I': 60, 'J': 10, 'K': 36, 'L': 16}
for k, v in widths.items(): ws.column_dimensions[k].width = v
for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
ws.freeze_panes = 'C2'
wb.save('kikfia-meta-catalog.xlsx')
print(len(rows), 'products')
